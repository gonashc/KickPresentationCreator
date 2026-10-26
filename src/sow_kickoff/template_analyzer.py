from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from .models import SlideInventory, TemplateInventory, TemplateTarget


MAX_CURRENT_TEXT = 600


def _clean(value: str) -> str:
    return " ".join((value or "").split())[:MAX_CURRENT_TEXT]


def _placeholder_name(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    try:
        return PP_PLACEHOLDER(shape.placeholder_format.type).name
    except Exception:
        return str(shape.placeholder_format.type)


def _slide_title(slide) -> str:
    title = slide.shapes.title
    if title is not None and getattr(title, "has_text_frame", False):
        return _clean(title.text)
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and _clean(shape.text):
            return _clean(shape.text)
    return ""


def analyze_template(data: bytes) -> TemplateInventory:
    if not data:
        raise ValueError("The PowerPoint template is empty.")
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise ValueError("The template is not a readable .pptx file.") from exc

    slides: list[SlideInventory] = []
    total_targets = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        targets: list[TemplateTarget] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                targets.append(
                    TemplateTarget(
                        kind="shape",
                        shape_id=shape.shape_id,
                        shape_name=shape.name,
                        current_text=_clean(shape.text),
                        placeholder_type=_placeholder_name(shape),
                    )
                )
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        targets.append(
                            TemplateTarget(
                                kind="table_cell",
                                shape_id=shape.shape_id,
                                shape_name=shape.name,
                                current_text=_clean(cell.text),
                                placeholder_type="TABLE_CELL",
                                row_index=row_index,
                                column_index=column_index,
                            )
                        )
        total_targets += len(targets)
        slides.append(SlideInventory(slide_number=slide_number, title=_slide_title(slide), targets=targets))

    if not slides:
        raise ValueError("The PowerPoint template contains no slides.")
    if total_targets == 0:
        raise ValueError("The PowerPoint template has no editable text or table cells.")

    return TemplateInventory(
        slide_count=len(slides),
        editable_target_count=total_targets,
        slides=slides,
    )

