from __future__ import annotations

import re
from io import BytesIO

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

from .models import DeckPlan, RenderResult


TOKEN_PATTERNS = {
    "client_name": [r"\{\{\s*CLIENT_NAME\s*\}\}", r"\[\s*Client Name\s*\]"],
    "project_name": [r"\{\{\s*PROJECT_NAME\s*\}\}", r"\[\s*Project Name\s*\]"],
    "kickoff_date": [r"\{\{\s*KICKOFF_DATE\s*\}\}", r"\[\s*Kickoff Date\s*\]"],
}


def _shape_by_id(slide, shape_id: int):
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    return None


def _replace_tokens(value: str, metadata: dict[str, str]) -> str:
    output = value
    for key, patterns in TOKEN_PATTERNS.items():
        replacement = metadata.get(key, "")
        for pattern in patterns:
            output = re.sub(pattern, replacement, output, flags=re.IGNORECASE)
    return output


def _write_text(text_frame, value: str) -> None:
    """Set text while retaining the shape and its inherited layout/theme styling."""
    lines = value.splitlines() if value else [""]
    text_frame.clear()
    text_frame.word_wrap = True
    try:
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass

    for index, raw_line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        line = raw_line.strip()
        if line.startswith(("•", "-", "–")):
            line = line[1:].strip()
        paragraph.text = line
        paragraph.level = 0


def _apply_metadata_tokens(presentation: Presentation, metadata: dict[str, str]) -> list[dict]:
    applied: list[dict] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                original = shape.text
                updated = _replace_tokens(original, metadata)
                if updated != original:
                    _write_text(shape.text_frame, updated)
                    applied.append(
                        {"slide": slide_number, "shape_id": shape.shape_id, "kind": "metadata_token"}
                    )
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        original = cell.text
                        updated = _replace_tokens(original, metadata)
                        if updated != original:
                            _write_text(cell.text_frame, updated)
                            applied.append(
                                {
                                    "slide": slide_number,
                                    "shape_id": shape.shape_id,
                                    "kind": "metadata_table_token",
                                    "row": row_index,
                                    "column": column_index,
                                }
                            )
    return applied


def render_presentation(*, template_bytes: bytes, plan: DeckPlan, metadata: dict[str, str]) -> RenderResult:
    try:
        presentation = Presentation(BytesIO(template_bytes))
    except Exception as exc:
        raise ValueError("The template could not be opened for rendering.") from exc

    applied = _apply_metadata_tokens(presentation, metadata)
    skipped: list[dict] = []
    seen: set[tuple] = set()

    for edit in plan.edits:
        key = (edit.slide_number, edit.target_kind, edit.shape_id, edit.row_index, edit.column_index)
        if key in seen:
            skipped.append({"edit": edit.model_dump(), "reason": "Duplicate target"})
            continue
        seen.add(key)

        if edit.slide_number < 1 or edit.slide_number > len(presentation.slides):
            skipped.append({"edit": edit.model_dump(), "reason": "Slide number does not exist"})
            continue

        slide = presentation.slides[edit.slide_number - 1]
        shape = _shape_by_id(slide, edit.shape_id)
        if shape is None:
            skipped.append({"edit": edit.model_dump(), "reason": "Shape ID does not exist on slide"})
            continue

        if edit.target_kind == "shape":
            if not getattr(shape, "has_text_frame", False):
                skipped.append({"edit": edit.model_dump(), "reason": "Shape has no text frame"})
                continue
            _write_text(shape.text_frame, edit.text)
        else:
            if not getattr(shape, "has_table", False):
                skipped.append({"edit": edit.model_dump(), "reason": "Shape is not a table"})
                continue
            if not (0 <= edit.row_index < len(shape.table.rows)):
                skipped.append({"edit": edit.model_dump(), "reason": "Table row does not exist"})
                continue
            if not (0 <= edit.column_index < len(shape.table.columns)):
                skipped.append({"edit": edit.model_dump(), "reason": "Table column does not exist"})
                continue
            _write_text(shape.table.cell(edit.row_index, edit.column_index).text_frame, edit.text)

        applied.append(
            {
                "slide": edit.slide_number,
                "shape_id": edit.shape_id,
                "kind": edit.target_kind,
                "source_references": edit.source_references,
            }
        )

    output = BytesIO()
    presentation.save(output)
    return RenderResult(
        presentation_bytes=output.getvalue(),
        applied_count=len(applied),
        skipped_count=len(skipped),
        applied_updates=applied,
        skipped_updates=skipped,
    )

