from __future__ import annotations

from io import BytesIO

from docx import Document

from .models import SourceDocument


MAX_SOURCE_CHARACTERS = 120_000


def extract_sow(data: bytes) -> SourceDocument:
    """Extract paragraphs and tables with stable, human-readable source references."""
    if not data:
        raise ValueError("The SoW file is empty.")

    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ValueError("The SoW is not a readable .docx file.") from exc

    sections: list[str] = []
    paragraph_count = 0
    for index, paragraph in enumerate(document.paragraphs, start=1):
        value = " ".join(paragraph.text.split())
        if not value:
            continue
        paragraph_count += 1
        style_name = paragraph.style.name if paragraph.style is not None else ""
        style_hint = f" ({style_name})" if style_name else ""
        sections.append(f"[Paragraph {index}{style_hint}] {value}")

    for table_index, table in enumerate(document.tables, start=1):
        sections.append(f"[Table {table_index}]")
        for row_index, row in enumerate(table.rows, start=1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            sections.append(f"[Table {table_index}, row {row_index}] " + " | ".join(cells))

    text = "\n".join(sections)
    if not text.strip():
        raise ValueError("The SoW does not contain readable paragraphs or tables.")

    truncated = len(text) > MAX_SOURCE_CHARACTERS
    if truncated:
        text = text[:MAX_SOURCE_CHARACTERS] + "\n[Document truncated by the application]"

    return SourceDocument(
        text=text,
        paragraph_count=paragraph_count,
        table_count=len(document.tables),
        truncated=truncated,
    )

