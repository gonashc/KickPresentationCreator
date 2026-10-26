from __future__ import annotations

from io import BytesIO

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from sow_kickoff.docx_reader import extract_sow
from sow_kickoff.models import DeckPlan, SlideEdit
from sow_kickoff.pptx_renderer import render_presentation
from sow_kickoff.template_analyzer import analyze_template


def _document_bytes() -> bytes:
    document = Document()
    document.add_heading("Statement of Work", level=1)
    document.add_paragraph("The engagement will deliver discovery and configuration services.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Date"
    table.cell(1, 0).text = "Kickoff"
    table.cell(1, 1).text = "October 20"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _presentation_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "{{CLIENT_NAME}} kickoff"
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    body.name = "OBJECTIVES"
    body.text = "Replace with objectives"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_extracts_docx_with_references():
    source = extract_sow(_document_bytes())
    assert source.paragraph_count == 2
    assert source.table_count == 1
    assert "[Paragraph 2" in source.text
    assert "[Table 1, row 2]" in source.text


def test_analyzes_and_renders_template():
    template = _presentation_bytes()
    inventory = analyze_template(template)
    body = next(target for target in inventory.slides[0].targets if target.shape_name == "OBJECTIVES")

    plan = DeckPlan(
        executive_summary="Discovery and configuration engagement.",
        edits=[
            SlideEdit(
                slide_number=1,
                target_kind="shape",
                shape_id=body.shape_id,
                row_index=-1,
                column_index=-1,
                text="Confirm requirements\nConfigure the solution",
                source_references=["Paragraph 2"],
            )
        ],
        warnings=[],
    )
    result = render_presentation(
        template_bytes=template,
        plan=plan,
        metadata={"client_name": "Acme", "project_name": "", "kickoff_date": "October 20, 2026"},
    )

    generated = Presentation(BytesIO(result.presentation_bytes))
    assert generated.slides[0].shapes.title.text == "Acme kickoff"
    rendered_body = next(shape for shape in generated.slides[0].shapes if shape.name == "OBJECTIVES")
    assert "Confirm requirements" in rendered_body.text
    assert result.applied_count == 2
    assert result.skipped_count == 0


def test_rejects_invalid_ai_target_without_corrupting_output():
    template = _presentation_bytes()
    plan = DeckPlan(
        executive_summary="",
        edits=[
            SlideEdit(
                slide_number=99,
                target_kind="shape",
                shape_id=12345,
                row_index=-1,
                column_index=-1,
                text="Unsafe",
                source_references=[],
            )
        ],
        warnings=[],
    )
    result = render_presentation(template_bytes=template, plan=plan, metadata={})
    assert result.skipped_count == 1
    Presentation(BytesIO(result.presentation_bytes))

